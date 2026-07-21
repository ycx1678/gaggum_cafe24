from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DESKTOP = Path.home() / "Desktop"
OUT_DIR = DESKTOP / "gaggum_option_simulator_final_client_pack_v79"
DOCX_PATH = OUT_DIR / "00_gaggum_option_simulator_manual_v79.docx"
PDF_PATH = OUT_DIR / "00_gaggum_option_simulator_manual_v79.pdf"
ZIP_PATH = DESKTOP / "gaggum_option_simulator_final_client_pack_v79.zip"

GENERATOR = ROOT / "tools" / "option_simulator_csv_generator.html"
SAMPLE_CSV = ROOT / "nd" / "images" / "simulator" / "products" / "1034" / "simulator.csv"

FONT_KO = "Apple SD Gothic Neo"
FONT_PDF = "ArialUnicode"
FONT_PDF_PATH = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT_KO
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    run.font.size = Pt(9.5)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F6F8")
    text = "\n".join(lines)
    set_cell_text(cell, text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
            run.font.size = Pt(9)
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str, fill: str = "EAF2F8") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = FONT_KO
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    run.font.size = Pt(10.5)
    paragraph.add_run("\n")
    body_run = paragraph.add_run(body)
    body_run.font.name = FONT_KO
    body_run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    body_run.font.size = Pt(10)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(header_cells[idx], "E8EEF5")
        set_cell_text(header_cells[idx], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def set_doc_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_KO
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [
        ("Title", 24, "111111"),
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = FONT_KO
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def build_docx() -> None:
    doc = Document()
    set_doc_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("가꿈 옵션 시뮬레이터 운영 매뉴얼")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("비개발자 담당자용 / 카페24 웹 FTP 기준 / v79")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(90, 90, 90)

    add_callout(
        doc,
        "한 줄 요약",
        "상품번호 폴더(`/web/upload/simulator/상품번호/`)에 PNG 이미지와 `simulator.csv`를 같이 올리면, 고객이 기존 카페24 옵션을 선택할 때 시뮬레이터 이미지가 자동으로 바뀝니다.",
        fill="F4F6F9",
    )

    doc.add_heading("1. 전달 파일 구성", level=1)
    add_table(
        doc,
        ["파일", "용도"],
        [
            ["00_gaggum_option_simulator_manual_v79.pdf / .docx", "지금 보고 있는 통합 운영 매뉴얼입니다."],
            ["01_csv_generator.html", "simulator.csv를 만드는 도구입니다. PC에서 더블클릭해서 사용합니다."],
            ["02_sample_simulator_csv_1034.csv", "래더 학원 책상 + ALO체어 세트 기준 샘플 CSV입니다."],
        ],
        widths=[2.65, 3.85],
    )

    doc.add_heading("2. 핵심 구조", level=1)
    add_table(
        doc,
        ["구분", "관리 위치", "설명"],
        [
            ["구매 옵션", "카페24 상품 옵션", "고객이 실제로 선택하고 주문에 반영되는 옵션입니다."],
            ["이미지", "웹 FTP", "상판, 프레임, 의자 등 시뮬레이터에 겹쳐 보일 PNG 파일입니다."],
            ["이미지 연결표", "simulator.csv", "어떤 옵션값일 때 어떤 이미지를 보여줄지 정하는 표입니다."],
        ],
        widths=[1.2, 1.75, 3.55],
    )
    add_callout(
        doc,
        "중요",
        "시뮬레이터는 옵션 버튼을 새로 만들지 않습니다. 옵션 버튼은 카페24 상품 옵션에서 관리하고, CSV는 그 선택값을 읽어 이미지만 바꿉니다.",
        fill="FFF6D9",
    )

    doc.add_heading("3. 가장 쉬운 운영 순서", level=1)
    add_numbers(
        doc,
        [
            "상품 상세 URL에서 product_no 값을 확인합니다.",
            "웹 FTP에 `/web/upload/simulator/상품번호/` 폴더를 만듭니다.",
            "PNG 이미지를 상품번호 폴더 안에 업로드합니다.",
            "`01_csv_generator.html`을 열어 옵션명과 이미지 경로를 입력합니다.",
            "`simulator.csv`를 다운로드합니다.",
            "`simulator.csv`를 상품번호 폴더 바로 아래에 업로드합니다.",
            "상품 상세페이지에서 옵션을 바꿔 이미지가 바뀌는지 확인합니다.",
        ],
    )

    doc.add_heading("4. 업로드 폴더 예시", level=1)
    add_code_block(
        doc,
        [
            "/web/upload/simulator/1034/",
            "  simulator.csv",
            "  top/",
            "    maple.png",
            "    white.png",
            "  frame/",
            "    white.png",
            "    black.png",
            "  seat/",
            "    black.png",
            "    beige.png",
        ],
    )

    doc.add_heading("5. CSV 생성기 사용법", level=1)
    add_numbers(
        doc,
        [
            "`01_csv_generator.html` 파일을 더블클릭합니다.",
            "상품번호를 입력합니다. 예: `1034`",
            "이미지가 바뀌는 옵션명만 입력합니다. 예: 책상 사이즈, 상판 컬러, 프레임 컬러, 의자 컬러",
            "이미지가 바뀌지 않는 옵션은 넣지 않습니다. 예: 배송 방식, 수량",
            "`image` 칸에는 상품번호 폴더 안의 상대경로를 입력합니다. 예: `top/maple.png`",
            "기본 조합은 `default` 칸에 `Y`를 넣습니다.",
            "`simulator.csv 다운로드` 버튼을 눌러 파일을 받습니다.",
        ],
    )

    doc.add_heading("6. 이미지 파일 규칙", level=1)
    add_table(
        doc,
        ["항목", "권장"],
        [
            ["파일 형식", "PNG"],
            ["배경", "투명 배경"],
            ["크기", "같은 상품 안의 모든 레이어는 같은 캔버스 크기"],
            ["파일명", "영문 소문자, 숫자, 하이픈 사용"],
            ["피해야 할 예", "상판 메이플.png, top white.png, seat(1).png"],
        ],
        widths=[1.45, 5.05],
    )

    doc.add_heading("7. 새 상품 추가", level=1)
    add_numbers(
        doc,
        [
            "카페24 관리자에서 상품 옵션을 먼저 등록합니다.",
            "상품번호를 확인합니다.",
            "웹 FTP에 `/web/upload/simulator/상품번호/` 폴더를 만듭니다.",
            "이미지 PNG를 준비해서 업로드합니다.",
            "CSV 생성기로 `simulator.csv`를 만듭니다.",
            "같은 상품번호 폴더에 `simulator.csv`를 업로드합니다.",
            "PC와 모바일 상세페이지에서 옵션 변경 테스트를 합니다.",
        ],
    )

    doc.add_heading("8. 옵션이 추가될 때", level=1)
    add_table(
        doc,
        ["상황", "처리 방법"],
        [
            ["새 옵션값 추가", "카페24 옵션값을 추가한 뒤 CSV에 해당 이미지 행을 추가합니다."],
            ["새 옵션군 추가 - 이미지 바뀜", "카페24 옵션군을 추가하고, CSV 생성기에서 옵션 컬럼과 이미지 행을 추가합니다."],
            ["새 옵션군 추가 - 이미지 안 바뀜", "카페24 옵션군만 추가합니다. CSV에는 넣지 않아도 됩니다."],
            ["이미지 파일명 변경", "웹 FTP의 파일명과 CSV의 image 값이 동일해야 합니다."],
        ],
        widths=[2.15, 4.35],
    )

    doc.add_heading("9. 확인 방법", level=1)
    doc.add_heading("PC", level=2)
    add_numbers(
        doc,
        [
            "상품 상세페이지를 엽니다.",
            "아래로 스크롤합니다.",
            "옵션 영역 왼쪽에 `옵션 시뮬레이터 보기` 버튼이 보이는지 확인합니다.",
            "버튼을 누르고 옵션을 변경합니다.",
            "이미지가 옵션에 맞게 바뀌면 정상입니다.",
        ],
    )
    doc.add_heading("모바일", level=2)
    add_numbers(
        doc,
        [
            "상품 상세페이지를 모바일 화면에서 엽니다.",
            "아래로 스크롤합니다.",
            "하단에 `옵션 시뮬레이터 보기` 버튼이 보이는지 확인합니다.",
            "버튼을 누르고 옵션을 변경합니다.",
            "이미지가 옵션에 맞게 바뀌면 정상입니다.",
        ],
    )

    doc.add_heading("10. 문제 해결", level=1)
    add_table(
        doc,
        ["증상", "확인할 것"],
        [
            ["시뮬레이터가 안 보임", "`/web/upload/simulator/상품번호/simulator.csv` 위치에 파일이 있는지 확인합니다."],
            ["이미지만 안 보임", "CSV의 `image` 경로와 실제 PNG 파일명이 같은지 확인합니다."],
            ["옵션을 바꿔도 이미지가 안 바뀜", "CSV 옵션명과 사이트 옵션명이 같은지 확인합니다."],
            ["특정 색상만 안 바뀜", "CSV 옵션값과 사이트 옵션값이 정확히 같은지 확인합니다."],
            ["이미지가 삐뚤어짐", "같은 상품의 모든 PNG 캔버스 크기와 위치가 같은지 확인합니다."],
            ["예전 이미지가 보임", "브라우저 강력 새로고침 후 다시 확인합니다."],
        ],
        widths=[2.0, 4.5],
    )

    doc.add_heading("11. 최종 체크리스트", level=1)
    add_bullets(
        doc,
        [
            "상품번호를 확인했다.",
            "`/web/upload/simulator/상품번호/` 폴더를 만들었다.",
            "PNG 이미지 파일명을 영문으로 정리했다.",
            "`simulator.csv` 파일명을 바꾸지 않았다.",
            "`simulator.csv`를 상품번호 폴더 바로 아래에 올렸다.",
            "PC에서 옵션 변경 시 이미지가 바뀌는지 확인했다.",
            "모바일에서 옵션 변경 시 이미지가 바뀌는지 확인했다.",
        ],
    )

    doc.save(DOCX_PATH)


def pdf_styles():
    if FONT_PDF_PATH.exists():
        pdfmetrics.registerFont(TTFont(FONT_PDF, str(FONT_PDF_PATH)))
    else:
        raise FileNotFoundError(FONT_PDF_PATH)

    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "BaseKo",
        parent=styles["BodyText"],
        fontName=FONT_PDF,
        fontSize=9.6,
        leading=13,
        spaceAfter=6,
        textColor=colors.HexColor("#111111"),
    )
    return {
        "title": ParagraphStyle("TitleKo", parent=base, fontSize=21, leading=26, alignment=TA_CENTER, spaceAfter=8, textColor=colors.HexColor("#111111")),
        "subtitle": ParagraphStyle("SubtitleKo", parent=base, fontSize=10.5, leading=14, alignment=TA_CENTER, textColor=colors.HexColor("#666666"), spaceAfter=16),
        "h1": ParagraphStyle("H1Ko", parent=base, fontSize=15, leading=19, spaceBefore=14, spaceAfter=8, textColor=colors.HexColor("#2E74B5")),
        "h2": ParagraphStyle("H2Ko", parent=base, fontSize=12, leading=15, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#1F4D78")),
        "body": base,
        "small": ParagraphStyle("SmallKo", parent=base, fontSize=8.6, leading=11, textColor=colors.HexColor("#444444")),
        "code": ParagraphStyle("CodeKo", parent=base, fontName=FONT_PDF, fontSize=8.7, leading=11, leftIndent=4, textColor=colors.HexColor("#111111")),
        "callout": ParagraphStyle("CalloutKo", parent=base, fontSize=10, leading=13, spaceAfter=0),
    }


def p(text: str, style):
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("`", "")
    )
    return Paragraph(text, style)


def pdf_table(headers: list[str], rows: list[list[str]], widths: list[float], st) -> Table:
    data = [[p(h, st["small"]) for h in headers]]
    for row in rows:
        data.append([p(cell, st["small"]) for cell in row])
    table = Table(data, colWidths=[w * inch for w in widths], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111111")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D0D8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def pdf_bullets(items: list[str], st, numbered: bool = False):
    bullet_type = "1" if numbered else "bullet"
    return ListFlowable(
        [ListItem(p(item, st["body"]), leftIndent=12) for item in items],
        bulletType=bullet_type,
        start="1",
        leftIndent=18,
        bulletFontName=FONT_PDF,
        bulletFontSize=9,
    )


def build_pdf() -> None:
    st = pdf_styles()
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title="가꿈 옵션 시뮬레이터 운영 매뉴얼",
    )
    story = []
    story.append(p("가꿈 옵션 시뮬레이터 운영 매뉴얼", st["title"]))
    story.append(p("비개발자 담당자용 / 카페24 웹 FTP 기준 / v79", st["subtitle"]))
    story.append(pdf_table(["핵심 요약"], [["상품번호 폴더(/web/upload/simulator/상품번호/)에 PNG 이미지와 simulator.csv를 같이 올리면, 고객이 기존 카페24 옵션을 선택할 때 시뮬레이터 이미지가 자동으로 바뀝니다."]], [7.1], st))
    story.append(Spacer(1, 10))

    story.append(p("1. 전달 파일 구성", st["h1"]))
    story.append(pdf_table(["파일", "용도"], [
        ["00_gaggum_option_simulator_manual_v79.pdf / .docx", "지금 보고 있는 통합 운영 매뉴얼입니다."],
        ["01_csv_generator.html", "simulator.csv를 만드는 도구입니다. PC에서 더블클릭해서 사용합니다."],
        ["02_sample_simulator_csv_1034.csv", "래더 학원 책상 + ALO체어 세트 기준 샘플 CSV입니다."],
    ], [3.0, 4.1], st))

    story.append(p("2. 핵심 구조", st["h1"]))
    story.append(pdf_table(["구분", "관리 위치", "설명"], [
        ["구매 옵션", "카페24 상품 옵션", "고객이 실제로 선택하고 주문에 반영되는 옵션입니다."],
        ["이미지", "웹 FTP", "상판, 프레임, 의자 등 시뮬레이터에 겹쳐 보일 PNG 파일입니다."],
        ["이미지 연결표", "simulator.csv", "어떤 옵션값일 때 어떤 이미지를 보여줄지 정하는 표입니다."],
    ], [1.25, 1.65, 4.2], st))
    story.append(pdf_table(["중요"], [["시뮬레이터는 옵션 버튼을 새로 만들지 않습니다. 옵션 버튼은 카페24 상품 옵션에서 관리하고, CSV는 그 선택값을 읽어 이미지만 바꿉니다."]], [7.1], st))

    story.append(p("3. 가장 쉬운 운영 순서", st["h1"]))
    story.append(pdf_bullets([
        "상품 상세 URL에서 product_no 값을 확인합니다.",
        "웹 FTP에 /web/upload/simulator/상품번호/ 폴더를 만듭니다.",
        "PNG 이미지를 상품번호 폴더 안에 업로드합니다.",
        "01_csv_generator.html을 열어 옵션명과 이미지 경로를 입력합니다.",
        "simulator.csv를 다운로드합니다.",
        "simulator.csv를 상품번호 폴더 바로 아래에 업로드합니다.",
        "상품 상세페이지에서 옵션을 바꿔 이미지가 바뀌는지 확인합니다.",
    ], st, numbered=True))

    story.append(p("4. 업로드 폴더 예시", st["h1"]))
    story.append(pdf_table(["예시 구조"], [[
        "/web/upload/simulator/1034/<br/>"
        "&nbsp;&nbsp;simulator.csv<br/>"
        "&nbsp;&nbsp;top/maple.png<br/>"
        "&nbsp;&nbsp;top/white.png<br/>"
        "&nbsp;&nbsp;frame/white.png<br/>"
        "&nbsp;&nbsp;frame/black.png<br/>"
        "&nbsp;&nbsp;seat/black.png<br/>"
        "&nbsp;&nbsp;seat/beige.png"
    ]], [7.1], st))

    story.append(p("5. CSV 생성기 사용법", st["h1"]))
    story.append(pdf_bullets([
        "01_csv_generator.html 파일을 더블클릭합니다.",
        "상품번호를 입력합니다. 예: 1034",
        "이미지가 바뀌는 옵션명만 입력합니다. 예: 책상 사이즈, 상판 컬러, 프레임 컬러, 의자 컬러",
        "이미지가 바뀌지 않는 옵션은 넣지 않습니다. 예: 배송 방식, 수량",
        "image 칸에는 상품번호 폴더 안의 상대경로를 입력합니다. 예: top/maple.png",
        "기본 조합은 default 칸에 Y를 넣습니다.",
        "simulator.csv 다운로드 버튼을 눌러 파일을 받습니다.",
    ], st, numbered=True))

    story.append(PageBreak())
    story.append(p("6. 이미지 파일 규칙", st["h1"]))
    story.append(pdf_table(["항목", "권장"], [
        ["파일 형식", "PNG"],
        ["배경", "투명 배경"],
        ["크기", "같은 상품 안의 모든 레이어는 같은 캔버스 크기"],
        ["파일명", "영문 소문자, 숫자, 하이픈 사용"],
        ["피해야 할 예", "상판 메이플.png, top white.png, seat(1).png"],
    ], [1.5, 5.6], st))

    story.append(p("7. 새 상품 추가", st["h1"]))
    story.append(pdf_bullets([
        "카페24 관리자에서 상품 옵션을 먼저 등록합니다.",
        "상품번호를 확인합니다.",
        "웹 FTP에 /web/upload/simulator/상품번호/ 폴더를 만듭니다.",
        "이미지 PNG를 준비해서 업로드합니다.",
        "CSV 생성기로 simulator.csv를 만듭니다.",
        "같은 상품번호 폴더에 simulator.csv를 업로드합니다.",
        "PC와 모바일 상세페이지에서 옵션 변경 테스트를 합니다.",
    ], st, numbered=True))

    story.append(p("8. 옵션이 추가될 때", st["h1"]))
    story.append(pdf_table(["상황", "처리 방법"], [
        ["새 옵션값 추가", "카페24 옵션값을 추가한 뒤 CSV에 해당 이미지 행을 추가합니다."],
        ["새 옵션군 추가 - 이미지 바뀜", "카페24 옵션군을 추가하고, CSV 생성기에서 옵션 컬럼과 이미지 행을 추가합니다."],
        ["새 옵션군 추가 - 이미지 안 바뀜", "카페24 옵션군만 추가합니다. CSV에는 넣지 않아도 됩니다."],
        ["이미지 파일명 변경", "웹 FTP의 파일명과 CSV의 image 값이 동일해야 합니다."],
    ], [2.5, 4.6], st))

    story.append(p("9. 확인 방법", st["h1"]))
    story.append(p("PC", st["h2"]))
    story.append(pdf_bullets([
        "상품 상세페이지를 엽니다.",
        "아래로 스크롤합니다.",
        "옵션 영역 왼쪽에 옵션 시뮬레이터 보기 버튼이 보이는지 확인합니다.",
        "버튼을 누르고 옵션을 변경합니다.",
        "이미지가 옵션에 맞게 바뀌면 정상입니다.",
    ], st, numbered=True))
    story.append(p("모바일", st["h2"]))
    story.append(pdf_bullets([
        "상품 상세페이지를 모바일 화면에서 엽니다.",
        "아래로 스크롤합니다.",
        "하단에 옵션 시뮬레이터 보기 버튼이 보이는지 확인합니다.",
        "버튼을 누르고 옵션을 변경합니다.",
        "이미지가 옵션에 맞게 바뀌면 정상입니다.",
    ], st, numbered=True))

    story.append(p("10. 문제 해결", st["h1"]))
    story.append(pdf_table(["증상", "확인할 것"], [
        ["시뮬레이터가 안 보임", "/web/upload/simulator/상품번호/simulator.csv 위치에 파일이 있는지 확인합니다."],
        ["이미지만 안 보임", "CSV의 image 경로와 실제 PNG 파일명이 같은지 확인합니다."],
        ["옵션을 바꿔도 이미지가 안 바뀜", "CSV 옵션명과 사이트 옵션명이 같은지 확인합니다."],
        ["특정 색상만 안 바뀜", "CSV 옵션값과 사이트 옵션값이 정확히 같은지 확인합니다."],
        ["이미지가 삐뚤어짐", "같은 상품의 모든 PNG 캔버스 크기와 위치가 같은지 확인합니다."],
        ["예전 이미지가 보임", "브라우저 강력 새로고침 후 다시 확인합니다."],
    ], [2.25, 4.85], st))

    story.append(p("11. 최종 체크리스트", st["h1"]))
    story.append(pdf_bullets([
        "상품번호를 확인했다.",
        "/web/upload/simulator/상품번호/ 폴더를 만들었다.",
        "PNG 이미지 파일명을 영문으로 정리했다.",
        "simulator.csv 파일명을 바꾸지 않았다.",
        "simulator.csv를 상품번호 폴더 바로 아래에 올렸다.",
        "PC에서 옵션 변경 시 이미지가 바뀌는지 확인했다.",
        "모바일에서 옵션 변경 시 이미지가 바뀌는지 확인했다.",
    ], st))

    doc.build(story)


def package_outputs() -> None:
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    build_docx()
    build_pdf()

    shutil.copy2(GENERATOR, OUT_DIR / "01_csv_generator.html")
    shutil.copy2(SAMPLE_CSV, OUT_DIR / "02_sample_simulator_csv_1034.csv")

    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(OUT_DIR.rglob("*")):
            archive.write(path, path.relative_to(DESKTOP))


if __name__ == "__main__":
    package_outputs()
